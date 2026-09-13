"""
test_report_exporter.py
=======================
Unit test suite for SAFEMETRIC Multi-Format Report Exporter (DOCX, CSV, PDF).
"""

import os
import sys
import csv
import docx

sys.path.insert(0, os.path.dirname(__file__))

from app.reports.report_exporter import export_report_docx, export_report_csv


def test_docx_export():
    """Verify DOCX generation and internal XML structure."""
    print("\n--- 1. Testing DOCX Report Generation ---")

    sample_inspection = {
        "inspection_id": "INS-TEST-1234",
        "product_name": "Premium Wheat Cookies",
        "created_at": "2026-09-12 11:30:00",
        "compliance_status": "COMPLIANT",
        "compliance_score": 92.5,
        "extracted_data": {
            "product_name": {"value": "Premium Wheat Cookies", "confidence": 96.0, "status": "Found"},
            "net_quantity": {"value": "200 g", "confidence": 98.0, "status": "Found"},
            "mrp": {"value": "Rs. 50.00", "confidence": 95.0, "status": "Found"},
            "manufacturer": {"value": "Bakeries India Ltd", "confidence": 94.0, "status": "Found"},
            "consumer_care": {"value": "care@bakeries.in", "confidence": 95.0, "status": "Found"},
            "font_analysis": {
                "overall_status": "PASS",
                "estimation_method": "user_calibrated_package_dimensions",
                "confidence": 0.95,
                "fields": {
                    "net_quantity": {
                        "field_label": "Net Quantity",
                        "char_height_mm": 2.4,
                        "min_required_mm": 1.0,
                        "status": "PASS"
                    },
                    "mrp": {
                        "field_label": "Maximum Retail Price (MRP)",
                        "char_height_mm": 2.1,
                        "min_required_mm": 1.0,
                        "status": "PASS"
                    }
                },
                "limitations": "Calibrated with package dimensions."
            },
            "placement_analysis": {
                "status": "PASS",
                "grouping": {
                    "status": "PASS",
                    "fields_analyzed": ["net_quantity", "mrp", "manufacturer"],
                    "explanation": "Declarations grouped on Principal Display Panel."
                },
                "clear_space": {
                    "status": "PASS",
                    "numeral_height_px": 25.0,
                    "required_vertical_clearance_px": 25.0,
                    "required_horizontal_clearance_px": 50.0,
                    "intrusions_detected": [],
                    "explanation": "Clear space zone verified."
                }
            }
        },
        "rule_results": [
            {"rule_id": "PCR-2011-R6-1-C-NET-QUANTITY", "status": "PASS"},
            {"rule_id": "PCR-2011-R6-1-E-MRP", "status": "PASS"}
        ],
        "violations": []
    }

    test_docx_path = os.path.join(os.path.dirname(__file__), "test_output_report.docx")
    try:
        out = export_report_docx(sample_inspection, test_docx_path, inspector_name="Officer Sharma")
        assert os.path.exists(out), "DOCX file was not created on disk."
        assert os.path.getsize(out) > 1000, "DOCX file is suspiciously small."

        # Load with python-docx to verify valid structure
        doc = docx.Document(out)
        text_content = "\n".join(p.text for p in doc.paragraphs)
        assert "GOVERNMENT OF INDIA" in text_content
        assert "STATUTORY PACKAGING COMPLIANCE INSPECTION REPORT" in text_content
        assert "Officer Verification & Sign-off Block" in text_content

        # Verify tables
        tables = doc.tables
        assert len(tables) >= 5, f"Expected at least 5 tables (metadata, exec, decl, font, placement), got {len(tables)}"

        print(f"  [PASS] DOCX report generated successfully ({os.path.getsize(out)} bytes, {len(tables)} tables).")
    finally:
        if os.path.exists(test_docx_path):
            os.remove(test_docx_path)


def test_csv_export():
    """Verify CSV generation and tabular column layout."""
    print("\n--- 2. Testing CSV Audit Export ---")

    sample_inspection = {
        "inspection_id": "INS-CSV-5678",
        "product_name": "Crunchy Chips",
        "created_at": "2026-09-12 11:45:00",
        "compliance_status": "NON-COMPLIANT",
        "compliance_score": 65.0,
        "extracted_data": {
            "net_quantity": {"value": "50 g", "confidence": 92.0, "status": "Found"},
            "mrp": {"value": "Rs. 20.00", "confidence": 95.0, "status": "Found"},
            "font_analysis": {
                "fields": {
                    "net_quantity": {"char_height_mm": 1.5, "min_required_mm": 1.0, "status": "PASS"},
                    "mrp": {"char_height_mm": 1.8, "min_required_mm": 1.0, "status": "PASS"}
                }
            },
            "placement_analysis": {
                "grouping": {"status": "PASS"},
                "clear_space": {"status": "FAIL"}
            }
        },
        "violations": [
            {"field": "net_quantity", "issue": "Clear space encroached", "severity": "HIGH"}
        ]
    }

    test_csv_path = os.path.join(os.path.dirname(__file__), "test_output_report.csv")
    try:
        out = export_report_csv(sample_inspection, test_csv_path, inspector_name="Officer Verma")
        assert os.path.exists(out), "CSV file was not created on disk."

        with open(out, mode="r", encoding="utf-8") as f:
            reader = csv.reader(f)
            rows = list(reader)

        assert len(rows) >= 3, f"Expected header + 2 field rows, got {len(rows)}"
        header = rows[0]
        assert "inspection_id" in header
        assert "overall_compliance_status" in header
        assert "font_char_height_mm" in header
        assert "clear_space_status" in header

        row1 = rows[1]
        assert row1[0] == "INS-CSV-5678"
        assert row1[3] == "Crunchy Chips"
        assert row1[4] == "NON-COMPLIANT"

        print(f"  [PASS] CSV export generated successfully ({len(rows)} rows, {len(header)} columns).")
    finally:
        if os.path.exists(test_csv_path):
            os.remove(test_csv_path)


if __name__ == "__main__":
    test_docx_export()
    test_csv_export()
    print("\n" + "=" * 70)
    print(" ALL REPORT EXPORTER TESTS (DOCX & CSV) PASSED (100%)!")
    print("=" * 70)
